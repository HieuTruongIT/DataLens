from io import BytesIO
from zipfile import ZipFile
import pandas as pd
import pytest
from openpyxl import Workbook
from datatriage.loading import load_data, sheet_names
from datatriage.validation import MAX_BYTES


def workbook():
    wb = Workbook()
    wb.active.title = "First"
    wb.active.append(["id", "id", None])
    wb.active.append(["001", 2, "=1+1"])
    wb.create_sheet("Second").append(["other"])
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


@pytest.mark.parametrize("delimiter", [",", ";", "\t", "|"])
def test_csv_preserves_leading_zero_and_literals(delimiter):
    result = load_data(f'id{delimiter}value\n001{delimiter}NA\n002{delimiter}\n'.encode(), "source.csv")
    assert result.frame.iloc[0].tolist() == ["001", "NA"]
    assert pd.isna(result.frame.iloc[1, 1])


def test_quoted_csv_utf8_and_bom():
    result = load_data('\ufeffname,note\nӘлия,"line 1, line 2"\n'.encode(), "x.csv")
    assert result.frame.iloc[0].tolist() == ["Әлия", "line 1, line 2"]


def test_headers_and_duplicates_are_not_lost():
    result = load_data(b"a,a,a [2],\n1,2,3,4\n", "x.csv")
    assert result.frame.columns.tolist() == ["a", "a [2]", "a [2] [2]", "Unnamed: 3"]
    assert len(result.notes) == 4


def test_header_only_is_valid_empty_dataset():
    assert load_data(b"value\n", "one.csv").frame.shape == (0, 1)


def test_xlsx_sheets_and_literal_formula():
    data = workbook()
    assert sheet_names(data) == ["First", "Second"]
    result = load_data(data, "x.xlsx", "First")
    assert result.frame.iloc[0].tolist() == ["001", 2, "=1+1"]
    assert load_data(data, "x.xlsx", "Second").frame.shape == (0, 1)


@pytest.mark.parametrize("data,name", [(b"", "x.csv"), (b"\xff", "x.csv"), (b"a,b\n1,2,3", "x.csv"),
                                      (b'a\n"unterminated', "x.csv"), (b"bad", "x.xlsx"),
                                      (b"a\n1", "x.exe"), (b"a\n\x00", "x.csv"), (b"\n", "x.csv")])
def test_bad_files_are_helpful(data, name):
    with pytest.raises(ValueError):
        load_data(data, name)


def test_missing_sheet():
    with pytest.raises(ValueError, match="sheet"):
        load_data(workbook(), "x.xlsx", "Nope")


def test_size_limit():
    with pytest.raises(ValueError, match="10 MB"):
        load_data(b"a" * (MAX_BYTES + 1), "x.csv")


def test_expansion_limit(monkeypatch):
    import datatriage.loading as module
    monkeypatch.setattr(module, "MAX_UNZIPPED_BYTES", 10)
    with pytest.raises(ValueError, match="expanded"):
        load_data(workbook(), "x.xlsx")


def test_macros_rejected():
    output = BytesIO()
    with ZipFile(output, "w") as z:
        z.writestr("xl/vbaProject.bin", b"synthetic")
    with pytest.raises(ValueError, match="Macro"):
        load_data(output.getvalue(), "x.xlsx")


def test_row_limit(monkeypatch):
    import datatriage.loading as module
    monkeypatch.setattr(module, "MAX_ROWS", 1)
    with pytest.raises(ValueError, match="limits"):
        load_data(b"a\n1\n2\n", "x.csv")


def test_streaming_xlsx_preserves_trailing_empty_columns_and_rows():
    from datatriage.reporting import xlsx_bytes
    original = pd.DataFrame({"name": ["a", None], "empty": [None, None]})
    result = load_data(xlsx_bytes(original), "export.xlsx").frame
    assert result.shape == (2, 2)
    assert result.iloc[0, 0] == "a"
    assert result["empty"].isna().all()
    assert result.iloc[1].isna().all()


def test_streaming_xlsx_wider_later_row_preserves_values():
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append(["a"])
    ws.append([1])
    ws.append([2, 3])
    output = BytesIO()
    wb.save(output)
    result = load_data(output.getvalue(), "export.xlsx")
    assert result.frame.shape == (2, 2)
    assert result.frame.columns.tolist() == ["a", "Unnamed: 1"]
    assert result.frame.iloc[1, 1] == 3
    assert pd.isna(result.frame.iloc[0, 1])
