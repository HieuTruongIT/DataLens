from io import BytesIO
import json
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from datatriage.models import Action, Settings, Rule
from datatriage.cleaning import clean
from datatriage.reporting import report_dict, json_bytes, csv_bytes, xlsx_bytes


def test_report_roundtrip_and_before_after():
    df = pd.DataFrame({"a": [" x ", "x", None], "n": [1, 1, np.inf]})
    result = clean(df, [Action("trim", "a"), Action("remove_duplicates")])
    report = json.loads(json_bytes(report_dict(df, result, Settings())))
    assert report["before"]["overview"]["rows"] == 3
    assert report["after"]["overview"]["rows"] == 2
    assert report["changes"]["rows_removed"] == 1
    assert len(report["actions_applied"]) == 2
    assert report["generated_at"].endswith("+00:00")


def test_report_omits_removed_column_rules_and_no_applied_result():
    df = pd.DataFrame({"a": ["bad"]})
    settings = Settings(rules=(Rule("a", "numeric"),))
    report = report_dict(df, clean(df, [Action("drop_column", "a")]), settings)
    assert report["after"]["score"]["overall"] is None
    assert report_dict(df, None, settings)["after"] is None


def test_rows_capped_but_counts_exact():
    df = pd.DataFrame({"a": [None] * 150})
    report = report_dict(df, None, Settings())
    missing = next(i for i in report["before"]["issues"] if i["kind"] == "missing")
    assert missing["affected_count"] == 150
    assert len(missing["rows"]) == 100
    assert missing["rows_truncated"]


def test_csv_formula_protection_is_explicit_and_copy_based():
    df = pd.DataFrame({"=header": ["=1+1", " @SUM(A1)", "-text", "safe"], "n": [-1, 2, 3, 4]})
    safe = csv_bytes(df).decode("utf-8-sig")
    raw = csv_bytes(df, False).decode("utf-8-sig")
    assert "'=header" in safe and "'=1+1" in safe and "' @SUM" in safe
    assert "'=1+1" not in raw
    assert df.iloc[0, 0] == "=1+1"
    assert ",-1" in safe


def test_xlsx_roundtrip_literal_formula_unicode_missing():
    df = pd.DataFrame({"=header": ["=1+1", "Әлия", None], "n": [1, np.inf, None]})
    wb = load_workbook(BytesIO(xlsx_bytes(df)), data_only=False)
    ws = wb.active
    assert ws["A1"].data_type == "s"
    assert ws["A2"].value == "=1+1" and ws["A2"].data_type == "s"
    assert ws["A3"].value == "Әлия"
    assert ws["B3"].value == "inf"


def test_excel_preserves_large_integer_as_text():
    wb = load_workbook(BytesIO(xlsx_bytes(pd.DataFrame({"id": [1234567890123456789]}))))
    assert wb.active["A2"].value == "1234567890123456789"


def test_excel_rejects_truncation_and_illegal_characters():
    import pytest
    for value in ["x" * 32768, "bad\x01text"]:
        with pytest.raises(ValueError, match="Excel cannot store"):
            xlsx_bytes(pd.DataFrame({"text": [value]}))
