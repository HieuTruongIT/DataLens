"""Bounded CSV/XLSX loading; preserve CSV values and never evaluate formulas."""
import csv
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile, BadZipFile
import pandas as pd
from openpyxl import load_workbook
from .models import LoadedData
from .validation import MAX_BYTES, MAX_ROWS, MAX_COLUMNS, MAX_CELLS, MAX_UNZIPPED_BYTES, validate_frame


def _check(data: bytes) -> None:
    if not data:
        raise ValueError("The file is empty. Choose a CSV or XLSX file containing a header.")
    if len(data) > MAX_BYTES:
        raise ValueError("File exceeds the 10 MB upload limit.")


def _workbook(data: bytes):
    _check(data)
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_UNZIPPED_BYTES:
                raise ValueError("Excel contents exceed the 50 MB expanded-size limit.")
            if any("vbaproject" in i.filename.lower() for i in archive.infolist()):
                raise ValueError("Macro-enabled workbooks are not supported.")
        return load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError, TypeError, SyntaxError) as exc:
        raise ValueError("Cannot read this XLSX workbook. Check that it is a valid Excel file.") from exc


def sheet_names(data: bytes) -> list[str]:
    try:
        wb = _workbook(data)
        try:
            return wb.sheetnames
        finally:
            wb.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Cannot read this XLSX workbook. It may be corrupted.") from exc


def _names(header: list) -> tuple[list[str], list[str]]:
    names, notes, used = [], [], set()
    for i, raw in enumerate(header):
        base = str(raw) if raw is not None and str(raw) != "" else f"Unnamed: {i}"
        name = base
        number = 2
        while name in used:
            name = f"{base} [{number}]"
            number += 1
        if name != base:
            notes.append(f'Duplicate column name "{base}" renamed to "{name}" on import.')
        if raw is None or str(raw) == "":
            notes.append(f'Blank header at column {i + 1} named "{name}" on import.')
        used.add(name)
        names.append(name)
    return names, notes


def load_data(data: bytes, filename: str, sheet: str | None = None, delimiter: str = "Auto") -> LoadedData:
    _check(data)
    ext = Path(filename).suffix.lower()
    rows, notes = [], []
    try:
        if ext == ".csv":
            text = data.decode("utf-8-sig")
            if "\x00" in text:
                raise ValueError("CSV contains binary/null content.")
            if delimiter == "Auto":
                try:
                    delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
                except csv.Error:
                    delimiter = ","
            if delimiter not in {",", ";", "\t", "|"}:
                raise ValueError("Choose a supported CSV delimiter.")
            iterator = csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True)
            header = next(iterator, [])
            for row in iterator:
                if not row:
                    row = [""] * len(header)
                if len(row) != len(header):
                    raise ValueError("CSV rows have inconsistent widths. Check the delimiter or quoting.")
                rows.append([pd.NA if v == "" else v for v in row])
                _bound(rows, header)
            notes.append("CSV values are preserved as text; only empty fields become missing. Type suggestions do not convert values.")
        elif ext == ".xlsx":
            wb = _workbook(data)
            try:
                if sheet is not None and sheet not in wb.sheetnames:
                    raise ValueError("The selected Excel sheet does not exist.")
                ws = wb[sheet or wb.sheetnames[0]]
                if ws.max_row and ws.max_row > MAX_ROWS + 1 or ws.max_column and ws.max_column > MAX_COLUMNS:
                    raise ValueError("Excel sheet exceeds row or column limits.")
                iterator = ws.iter_rows(values_only=True)
                header = list(next(iterator, ()))
                for row in iterator:
                    row = list(row)
                    if len(row) > len(header):
                        header.extend([None] * (len(row) - len(header)))
                    rows.append(row)
                    _bound(rows, header)
                # Streaming XLSX files may omit dimensions and trailing empty cells.
                rows = [row + [None] * (len(header) - len(row)) for row in rows]
            finally:
                wb.close()
            notes.append("Excel formulas are imported as literal text, never evaluated. Styling and merged-cell structure are not retained.")
        else:
            raise ValueError("Unsupported file type. Upload .csv or .xlsx.")
    except UnicodeDecodeError as exc:
        raise ValueError("CSV must use UTF-8 encoding. Re-save your file as UTF-8 CSV.") from exc
    except csv.Error as exc:
        raise ValueError("CSV could not be parsed. Check the delimiter and quoted fields.") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("File could not be read. It may be corrupted or unsupported.") from exc
    if not header:
        raise ValueError("No header found. Add column names before importing.")
    columns, header_notes = _names(list(header))
    df = pd.DataFrame(rows, columns=columns)
    validate_frame(df)
    return LoadedData(df, header_notes + notes)


def _bound(rows: list, header: list) -> None:
    if len(rows) > MAX_ROWS or len(header) > MAX_COLUMNS or len(rows) * len(header) > MAX_CELLS:
        raise ValueError("Dataset exceeds limits: 50,000 rows, 200 columns or 2 million cells.")
